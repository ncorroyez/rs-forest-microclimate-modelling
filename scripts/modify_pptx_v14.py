"""
Corroyez_MEB2026_v13.pptx → v14.pptx.

Reinsert the regenerated forward-selection figure (cleaner annotations:
italic r in bold, RMSE, MAE — no R², no caption, bigger metrics, simplified
axis labels).
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

SRC = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v13.pptx")
DST = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v14.pptx")
FIG_FORWARD = Path("/home/corroyez/Documents/z_Example_rmusica_31012025/"
                    "outputs/figs_MEB2026_final/fig_forward_selection_HOBO.png")

prs = Presentation(SRC)


def set_notes(slide, text):
    nt = slide.notes_slide.notes_text_frame; nt.clear()
    for i, line in enumerate(text.split("\n")):
        p = nt.paragraphs[0] if i == 0 else nt.add_paragraph()
        r = p.add_run(); r.text = line; r.font.size = Pt(11)


# Slide 10 : replace forward-selection figure
slide10 = prs.slides[9]
pics = [(j, sh) for j, sh in enumerate(slide10.shapes) if sh.shape_type == 13]
pics_sorted = sorted(pics, key=lambda p: -(p[1].width * p[1].height))
if pics_sorted and FIG_FORWARD.exists():
    big_j = pics_sorted[0][0]
    old = slide10.shapes[big_j]
    left, top, width, height = old.left, old.top, old.width, old.height
    sp = old._element; sp.getparent().remove(sp)
    slide10.shapes.add_picture(str(FIG_FORWARD), left, top,
                                  width=width, height=height)
    print("Replaced forward-selection figure on slide 10")

# Update slide 10 notes — drop R² mention, keep r + RMSE + MAE
set_notes(prs.slides[9], (
"From theory to data. For each of the 53 HOBO sensors we computed the "
"observed ΔTmax on the x-axis and the simulated ΔTmax from MuSICA at "
"each forward-selection step on the y-axis. Same convention on both axes: "
"negative means the canopy buffers. We start from the uniform baseline, "
"then add traits one at a time in the LOO order: fCover, LAI, Hmax, "
"finally LAD = REF.\n"
"\n"
"Reading: Pearson r improves from −0.4 with the baseline to +0.53 with "
"full structure restored. RMSE and MAE stay around 2 °C across steps — "
"this is the glass-ceiling: simulated buffering spans ~2 °C while the "
"observed range is ~7 °C. So MuSICA gets part of the ordering, not the "
"magnitude.\n"
"\n"
"Notable: r is negative for the first two steps (fCover, fCover+LAI), "
"then flips to positive once Hmax is added. So Hmax is critical for "
"recovering the plot ordering — even though it only ranks third in the "
"LOO MAE. The MAE-rank and the ordering-rank are not the same question."
))

# Slide 8 notes — keep r/RMSE/MAE (drop R² mention)
set_notes(prs.slides[7], (
"Two-step framework.\n"
"\n"
"Step 1 — Attribution. We run 5 simulations per plot: the full real "
"canopy (REF), and four Leave-One-Out variants where one trait is "
"replaced by its baseline. Baselines, shown bottom-left: LAI and Hmax "
"replaced by the cLHS spatial mean; fCover by the same mean floored at "
"0.5; LAD by a uniform vertical profile. Δ_v = Tmax(LOO_v) − Tmax(REF). "
"Positive Δ_v means removing v warmed the canopy, so v contributes to "
"buffering.\n"
"\n"
"Step 2 — HOBO validation. We start from the uniform baseline and add "
"traits one at a time in the LOO order, then quantify how well the "
"simulated ΔTmax matches the observed one at each step. Pearson r for "
"correlation, RMSE and MAE for error magnitude.\n"
"\n"
"Why LOO and not LVA / Shapley? LVA tests each trait alone vs. the null "
"— unrealistic, misses interactions. Shapley uses all 16 trait "
"combinations but ends up close to LOO here. LOO is the simplest, most "
"transparent attribution — and matches the forward-selection narrative.\n"
"\n"
"Convention: ΔTmax = Tmicro − Tmacro, so negative means the canopy buffers."
))

prs.save(DST)
print(f"Saved {DST}")
