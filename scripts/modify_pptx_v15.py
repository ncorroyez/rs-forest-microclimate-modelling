"""
Corroyez_MEB2026_v14.pptx → v15.pptx.

Just re-insert the regenerated forward-selection figure
(r/RMSE/MAE same size/font/color, axes 'Field' and 'Simulated').
"""
from pathlib import Path
from pptx import Presentation

SRC = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v14.pptx")
DST = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v15.pptx")
FIG_FORWARD = Path("/home/corroyez/Documents/z_Example_rmusica_31012025/"
                    "outputs/figs_MEB2026_final/fig_forward_selection_HOBO.png")

prs = Presentation(SRC)
slide10 = prs.slides[9]
pics = [(j, sh) for j, sh in enumerate(slide10.shapes) if sh.shape_type == 13]
pics_sorted = sorted(pics, key=lambda p: -(p[1].width * p[1].height))
if pics_sorted and FIG_FORWARD.exists():
    big_j = pics_sorted[0][0]
    old = slide10.shapes[big_j]
    left, top, width, height = old.left, old.top, old.width, old.height
    sp = old._element; sp.getparent().remove(sp)
    slide10.shapes.add_picture(str(FIG_FORWARD), left, top,
                                  width=width, height=height)
    print("Replaced forward-selection figure on slide 10")
prs.save(DST)
print(f"Saved {DST}")
