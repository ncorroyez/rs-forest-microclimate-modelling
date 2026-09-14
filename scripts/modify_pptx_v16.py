"""
Corroyez_MEB2026_v15.pptx → v16.pptx.

Fix fCover baseline : baseline = 1 (closed canopy), not the spatial mean.
- Re-insert the regenerated baselines histogram figure on slide 8
- Update slide 8 text + notes accordingly
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

SRC = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v15.pptx")
DST = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v16.pptx")
FIG_BASE = Path("/home/corroyez/Documents/z_Example_rmusica_31012025/"
                 "outputs/figs_MEB2026_final/fig_baselines_distributions.png")

prs = Presentation(SRC)


def set_paragraph_text(shape, new_lines, font_size=None, bold=False):
    if not shape.has_text_frame: return
    tf = shape.text_frame; tf.clear()
    for i, line in enumerate(new_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run(); run.text = line
        if font_size: run.font.size = Pt(font_size)
        if bold:      run.font.bold = True


def set_notes(slide, text):
    nt = slide.notes_slide.notes_text_frame; nt.clear()
    for i, line in enumerate(text.split("\n")):
        p = nt.paragraphs[0] if i == 0 else nt.add_paragraph()
        r = p.add_run(); r.text = line; r.font.size = Pt(11)


# Replace baselines histogram on slide 8 (left column, ~4.7x1.3 in @ 0.2,2.9)
slide8 = prs.slides[7]
from pptx.util import Inches
for shape in list(slide8.shapes):
    if shape.shape_type != 13: continue
    # baselines fig was at (0.2, 2.9) with size ~4.7 x 1.3
    if abs(shape.left - Inches(0.2)) < Inches(0.1) \
       and abs(shape.top - Inches(2.9)) < Inches(0.2):
        left, top, w, h = shape.left, shape.top, shape.width, shape.height
        sp = shape.element; sp.getparent().remove(sp)
        slide8.shapes.add_picture(str(FIG_BASE), left, top, width=w, height=h)
        print("Replaced baselines fig on slide 8")
        break

# Update baselines text on slide 8
for shape in slide8.shapes:
    if shape.has_text_frame and "Baselines" in shape.text_frame.text:
        set_paragraph_text(shape, [
            "Factorial design: 4 LiDAR traits × {real, baseline} = 16 configurations",
            "",
            "Baselines — what replaces each trait when removed:",
            "  LAI / H_max → cLHS spatial mean",
            "  fCover      → 1 (closed canopy)",
            "  LAD         → uniform vertical profile",
        ], font_size=12)
        # add bold on heading
        para0 = shape.text_frame.paragraphs[0]
        if para0.runs: para0.runs[0].font.bold = True
        para2 = shape.text_frame.paragraphs[2]
        if para2.runs: para2.runs[0].font.bold = True
        break

# Update notes slide 8 — fCover baseline = 1
set_notes(prs.slides[7], (
"Two-step framework.\n"
"\n"
"Step 1 — Attribution. We run 5 simulations per plot: the full real "
"canopy (REF), and four Leave-One-Out variants where one trait is "
"replaced by its baseline. Baselines, shown bottom-left: LAI and Hmax "
"replaced by the cLHS spatial mean; fCover set to 1 (closed canopy — "
"three of four profiles already have fCover above 0.98); LAD replaced "
"by a uniform vertical profile that preserves the plot's own LAI and "
"Hmax. Δ_v = Tmax(LOO_v) − Tmax(REF). Positive Δ_v means removing v "
"warmed the canopy, so v contributes to buffering.\n"
"\n"
"Step 2 — HOBO validation. We start from the uniform baseline and add "
"traits one at a time in the LOO order, then quantify how well the "
"simulated ΔTmax matches the observed one at each step. Pearson r for "
"correlation, RMSE and MAE for error magnitude.\n"
"\n"
"Why LOO and not LVA / Shapley? LVA tests each trait alone vs. the null "
"— unrealistic, misses interactions. Shapley uses all 16 trait "
"combinations but ends up close to LOO here. LOO is the simplest, most "
"transparent attribution — and matches the forward-selection narrative.\n"
"\n"
"Convention: ΔTmax = Tmicro − Tmacro, so negative means the canopy buffers."
))

prs.save(DST)
print(f"Saved {DST}")
