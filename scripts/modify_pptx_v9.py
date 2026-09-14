"""
Corroyez_MEB2026_v8.pptx → v9.pptx.

Fix leftover 'baseline' references on slide 4 speaker notes (the histograms
are now on slide 8, not 4). Update slide 8 notes to mention them.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt

SRC = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v8.pptx")
DST = Path("/home/corroyez/Downloads/1Papers/files/Corroyez_MEB2026_v9.pptx")

prs = Presentation(SRC)


def set_notes(slide, text):
    nt = slide.notes_slide.notes_text_frame; nt.clear()
    for i, line in enumerate(text.split("\n")):
        p = nt.paragraphs[0] if i == 0 else nt.add_paragraph()
        r = p.add_run(); r.text = line; r.font.size = Pt(11)


# Slide 4 — MuSICA : keep only the tool description, no baselines / factorial
set_notes(prs.slides[3], (
"MuSICA — Multilayer Simulator of the Interactions between a Canopy and the "
"Atmosphere. A mechanistic model of energy / water / carbon exchange between "
"canopy and atmosphere, dividing the canopy into vertical layers (ideal for "
"microclimate stratification). Inputs: canopy structure (Hmax, LAI / LAD "
"profile, fCover), soil, and ERA5 macroclimate forcing. Output: hourly air "
"temperature at every layer, including 1 m above the ground."
))

# Slide 8 — Methods : add factorial + baselines context to the notes
set_notes(prs.slides[7], (
"Two-step procedure.\n"
"\n"
"Design — a 16-configuration factorial: each of the four LiDAR-derived "
"traits is independently set to its real plot value or to a baseline (see "
"the histograms left). LAI and H_max use the cLHS spatial mean, fCover the "
"same mean floored at 0.5 (microsite-level cover saturation), LAD a uniform "
"vertical profile.\n"
"\n"
"Step 1 — Theoretical attribution with Leave-One-Out. Start from the full "
"canopy with all four traits real, remove each trait one at a time, and "
"look at the resulting ΔTmax. Δ_v = T_max(LOO_v) − T_max(REF). Positive "
"Δ_v means removing v warmed the canopy, so v contributes to buffering.\n"
"\n"
"Step 2 — HOBO validation by forward selection. Start from the uniform "
"baseline and add traits one at a time in the LOO order; quantify how well "
"the simulated ΔTmax matches the observed ΔTmax at each step (Spearman ρ "
"for rank agreement, linear regression for magnitude).\n"
"\n"
"Why LOO over LVA / Shapley? LVA tests each trait alone vs. the null and "
"misses interactions. Shapley adds desirable axioms but behaves similarly "
"to LOO here. LOO is the simplest, most transparent, baseline-anchored "
"attribution — and matches the forward-selection narrative best."
))

prs.save(DST)
print(f"Saved {DST}")

# Sanity check
prs2 = Presentation(DST)
print()
for i in [3, 7]:
    notes = prs2.slides[i].notes_slide.notes_text_frame.text
    has_base = "baseline" in notes.lower()
    print(f"Slide {i+1} notes contain 'baseline'? {has_base}")
